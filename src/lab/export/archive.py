import io
import json
import os
import shutil
import tempfile
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from uuid import UUID, uuid4
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from lab.config import ROOT
from lab.db import serial, many, one, insert, transaction, value
from lab.domain.research import canonical, digest, input_hash
from lab.export.analysis import summarize, selected_runs
from lab.storage.client import Storage

SHEETS=("요약","논문_분기점","후보_아이디어","모델_선택","유용성_평가","일치도_평가","사람_검토","오류_미완료","실행_조건","긴텍스트_목록")
TABLES=("papers","checkpoints","checkpoint_revisions","model_profiles","protocol_versions","plans","runs",
        "jobs","attempts","ideas","selections","usefulness_ratings","agreement_ratings",
        "review_batches","review_assignments","human_reviews","audit_events")
ASSET_KINDS=("source_pdf","prepared_pdf","page_png","request_json","response_json","raw_text")

def safe_cell(cell,value):
    if isinstance(value,(dict,list)):value=canonical(value)
    if value is None:return
    if isinstance(value,(UUID,datetime)):value=str(value)
    cell.value=value
    if isinstance(value,str):cell.data_type="s"

def write_excel(rows,reviews,target,long_texts=None,common=True):
    long_texts=long_texts if long_texts is not None else {}
    wb=Workbook();wb.remove(wb.active)
    for name in SHEETS:wb.create_sheet(name)
    counters=defaultdict(int)
    def table(name,headers,data):
        ws=wb[name]
        for col,h in enumerate(headers,1):safe_cell(ws.cell(1,col),h)
        for n,row in enumerate(data,2):
            for col,item in enumerate(row,1):
                text=canonical(item) if isinstance(item,(list,dict)) else item
                if isinstance(text,str) and len(text)>30000:
                    key="long-text/"+str(len(long_texts)+1).zfill(6)+".txt"
                    long_texts[key]=text.encode()
                    counters["long"]+=1
                    wb["긴텍스트_목록"].append([key,name,headers[col-1],len(text),key])
                    text="[긴 본문: "+key+"]"
                safe_cell(ws.cell(n,col),text)
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill("solid",fgColor="E8EEF8");cell.font=Font(size=10,bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font=Font(size=10)
                if cell.data_type=="n":cell.number_format="0.###"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(60,max(12,max(len(str(c.value or "")) for c in col[:30])+2))
    summary=[]
    groups=sorted({r["group_hash"] for r in rows})
    for group in groups:
        agg=summarize(rows,common,group)
        for provider,values in agg["summary"].items():
            for metric,result in values.items():
                summary.append([group,provider,metric,result["value"],result["eligible_runs"],result["total_runs"],
                                result["papers"],result["excluded_runs"],agg["analysis_hash"]])
    table("요약",["조건그룹","생성제공사","지표","값","적격실행","전체실행","논문수","제외실행","분석hash"],summary)
    table("논문_분기점",["논문ID","코드","제목","분기점ID","버전ID","입력방식","계획ID"],
          [[r["paper_id"],r["paper_code"],r["paper_title"],r["checkpoint_code"],r["checkpoint_revision_id"],r["input_mode"],r["plan_id"]] for r in rows])
    table("후보_아이디어",["run_id","idea_id","후보","제목","핵심아이디어","도움이유","다음단계","가정","미검증","선택됨"],
          [[r["id"],i["id"],i["candidate_id"],i["title"],i["idea_md"],i["why_helpful_md"],i["next_steps"],i["assumptions"],i["unverified_points"],i["selected"]] for r in rows for i in r["ideas"]])
    table("모델_선택",["run_id","선택후보","근거","위험","attempt_id"],
          [[r["id"],r["selection"]["selected_ordinal"],r["selection"]["rationale_md"],r["selection"]["risks_md"],r["selection"]["source_attempt_id"]] for r in rows if r["selection"]])
    table("유용성_평가",["run_id","idea_id","평가자","점수","판단","근거","지적","판단불가이유","attempt_id"],
          [[r["id"],i["id"],v["judge_provider"],v["score"],v["validity"],v["rationale_md"],v["concerns"],v["unable_reason"],v["source_attempt_id"]] for r in rows for i in r["ideas"] for v in i["usefulness_ratings"]])
    table("일치도_평가",["run_id","idea_id","평가자","라벨","근거","대응내용","판단불가이유","attempt_id"],
          [[r["id"],i["id"],v["judge_provider"],v["label"],v["rationale_md"],v["corresponding_elements"],v["unable_reason"],v["source_attempt_id"]] for r in rows for i in r["ideas"] for v in i["agreement_ratings"]])
    table("사람_검토",["배정ID","검토자ID","단계","수정번호","점수","판단","일치도","근거","미판정이유","수정이유"],
          [[v["assignment_id"],v["created_by"],v["phase"],v["revision_no"],v["score"],v["validity"],v["agreement_label"],v["rationale_md"],v["unable_reason"],v["correction_reason"]] for v in reviews])
    table("오류_미완료",["run_id","job_id","단계","제공사","상태","오류유형","오류","최근변경"],
          [[r["id"],j["id"],j["phase"],j["provider"],j["status"],j["error_class"],j["error_message"],j["updated_at"]] for r in rows for j in r["jobs"] if j["status"]!="completed"])
    table("실행_조건",["run_id","구분","경로","입력방식","생성제공사","모델","모델프로필hash","평가자조합","입력hash","반복","사람보정","고정설정hash","UTC","한국시간"],
          [[r["id"],r["cohort"],r["transport"],r["input_mode"],r["generator_provider"],r["generator_model"]["model_id"],digest(r["generator_model"]),
            r["judge_models"],r["request_fingerprint"],r["repeat_no"],r["human_corrected"],r["frozen_sha256"],str(r["created_at"]),
            r["created_at"].astimezone(__import__("zoneinfo").ZoneInfo("Asia/Seoul")).isoformat()] for r in rows])
    index=wb["긴텍스트_목록"]
    index.insert_rows(1)
    for c,h in enumerate(("ID","시트","필드","전체길이","원문경로"),1):safe_cell(index.cell(1,c),h)
    index.freeze_panes="A2";index.auto_filter.ref=index.dimensions
    for cell in index[1]:cell.fill=PatternFill("solid",fgColor="E8EEF8");cell.font=Font(size=10,bold=True)
    wb.save(target)
    return counters["long"]

def gather(conn,user,filters,include_source=True,backup=False):
    ws=user["workspace_id"]
    rows=selected_runs(conn,user,filters)
    if backup:
        papers=many(conn,"select * from lab.papers where workspace_id=%s",(ws,))
        data={table:many(conn,"select * from lab."+table+" where workspace_id=%s",(ws,)) for table in TABLES}
    else:
        plan_ids={str(r["plan_id"]) for r in rows};run_ids={str(r["id"]) for r in rows}
        paper_ids={str(r["paper_id"]) for r in rows}
        data={}
        for table in TABLES:
            allrows=many(conn,"select * from lab."+table+" where workspace_id=%s",(ws,))
            if table=="papers":keep=[x for x in allrows if str(x["id"]) in paper_ids]
            elif table=="plans":keep=[x for x in allrows if str(x["id"]) in plan_ids]
            elif table=="runs":keep=[x for x in allrows if str(x["id"]) in run_ids]
            elif table in ("ideas","selections","usefulness_ratings","agreement_ratings","jobs","attempts"):
                jobs={str(j["id"]) for r in rows for j in r["jobs"]}
                keep=[x for x in allrows if str(x.get("run_id")) in run_ids or str(x.get("job_id")) in jobs]
            else:keep=allrows
            data[table]=keep
        checkpoints={str(r["checkpoint_revision_id"]) for r in rows}
        data["checkpoint_revisions"]=[r for r in data["checkpoint_revisions"] if str(r["id"]) in checkpoints]
        cpids={str(r["checkpoint_id"]) for r in data["checkpoint_revisions"]}
        data["checkpoints"]=[r for r in data["checkpoints"] if str(r["id"]) in cpids]
        protocol_ids={str(r["protocol_id"]) for r in data["plans"]}
        data["protocol_versions"]=[r for r in data["protocol_versions"] if str(r["id"]) in protocol_ids]
        profiles={str(v.get("id")) for r in rows for v in r["frozen_config"]["models"].values()}
        data["model_profiles"]=[r for r in data["model_profiles"] if str(r["id"]) in profiles]
        idea_ids={str(r["id"]) for r in data["ideas"]}
        data["review_assignments"]=[r for r in data["review_assignments"] if str(r["idea_id"]) in idea_ids]
        assignment_ids={str(r["id"]) for r in data["review_assignments"]}
        batch_ids={str(r["review_batch_id"]) for r in data["review_assignments"]}
        data["review_batches"]=[r for r in data["review_batches"] if str(r["id"]) in batch_ids]
        data["human_reviews"]=[r for r in data["human_reviews"] if str(r["assignment_id"]) in assignment_ids]
        object_ids=plan_ids|run_ids|paper_ids|checkpoints|cpids|idea_ids|assignment_ids|batch_ids
        data["audit_events"]=[r for r in data["audit_events"] if str(r.get("object_id")) in object_ids]
        papers=data["papers"]
    assets=many(conn,"select * from lab.assets where workspace_id=%s and kind=any(%s)",(ws,list(ASSET_KINDS)))
    if not include_source:assets=[a for a in assets if a["kind"]!="source_pdf"]
    # Include only assets referenced by the selected dataset, unless creating the full research backup.
    if not backup:
        mentioned=canonical(data)
        assets=[a for a in assets if str(a["id"]) in mentioned]
    if not include_source:
        for table in ("papers","checkpoint_revisions"):
            for row in data[table]:row["source_asset_id"]=None
    return data,assets,rows

def write_archive(user,filters,path,include_source=True,external=False,backup=False):
    store=Storage()
    with transaction() as conn:
        conn.execute("set transaction isolation level repeatable read, read only")
        data,assets,rows=gather(conn,user,filters,include_source and not external,backup)
        reviews=data["human_reviews"]
        expected=filters.get("analysis_hash")
        if expected and summarize(rows,filters.get("common",True),filters.get("group_hash") or None)["analysis_hash"]!=expected:
            raise ValueError("ANALYSIS_SNAPSHOT_CHANGED")
    if external:
        def scrub(obj):
            if isinstance(obj,dict):
                return {k:scrub(v) for k,v in obj.items() if k not in
                        ("email","created_by","user_id","reviewed_by","actor_id","reviewer_id","display_name")}
            if isinstance(obj,list):return [scrub(x) for x in obj]
            return obj
        data=scrub(data)
    manifest={"schema_version":"2.0","workspace_id":str(user["workspace_id"]),"created_at":datetime.now(timezone.utc).isoformat(),
              "filters":filters,"backup":backup,"external":external,"source_included":include_source and not external,
              "table_counts":{k:len(v) for k,v in data.items()},"files":{},"assets":[]}
    size=sum(a["byte_size"] or 0 for a in assets)
    if shutil.disk_usage(Path(path).parent).free < size*2+64*1024*1024:
        raise OSError(28,"Not enough temporary space")
    def add(z,name,blob):
        z.writestr(name,blob);manifest["files"][name]={"sha256":digest(blob),"bytes":len(blob)}
    with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED,allowZip64=True) as z:
        for table,records in data.items():
            add(z,"data/"+table+".jsonl",b"".join((canonical(r)+"\n").encode() for r in records))
        long={}
        excel=io.BytesIO()
        write_excel(rows,reviews,excel,long,common=filters.get("common",True))
        add(z,"results.xlsx",excel.getvalue())
        for name,blob in long.items():add(z,name,blob)
        for a in assets:
            if a["status"]!="verified":
                manifest["assets"].append({**a,"archive_path":None})
                continue
            content=store.get(a)
            if digest(content)!=a["sha256"]:raise ValueError("ASSET_HASH_MISMATCH")
            prefix={"source_pdf":"papers","prepared_pdf":"prepared","request_json":"requests",
                    "response_json":"responses","raw_text":"responses","page_png":"prepared"}[a["kind"]]
            name=prefix+"/"+str(a["id"])+"/"+Path(a["original_name"]).name
            add(z,name,content);manifest["assets"].append({**a,"archive_path":name})
        for p in (ROOT/"contracts").rglob("*.json"):
            if p.name=="deployment-config.template.json" or p.name=="env-contract.json":continue
            add(z,"protocols/"+str(p.relative_to(ROOT/"contracts")),p.read_bytes())
        for group in {r["group_hash"] for r in rows}:
            add(z,"analysis/"+group+".json",canonical(summarize(rows,filters.get("common",True),group)).encode())
        z.writestr("manifest.json",canonical(manifest))
    return manifest

def verify_archive(path):
    with zipfile.ZipFile(path) as z:
        infos=z.infolist()
        if len(infos)>100000:raise ValueError("ARCHIVE_TOO_MANY_FILES")
        total=0
        for info in infos:
            name=PurePosixPath(info.filename)
            if name.is_absolute() or ".." in name.parts or "\\" in info.filename or ":" in name.parts[0]:
                raise ValueError("UNSAFE_ARCHIVE_PATH")
            if (info.external_attr>>16)&0o170000==0o120000:raise ValueError("ARCHIVE_SYMLINK_FORBIDDEN")
            total+=info.file_size
            if info.compress_size and info.file_size/info.compress_size>1000:raise ValueError("ARCHIVE_COMPRESSION_RATIO")
        if total>20*1024**3:raise ValueError("ARCHIVE_TOO_LARGE")
        manifest=json.loads(z.read("manifest.json"))
        if manifest["schema_version"]!="2.0":raise ValueError("UNSUPPORTED_ARCHIVE_VERSION")
        for name,meta in manifest["files"].items():
            if name not in z.namelist():raise ValueError("ARCHIVE_FILE_MISSING")
            content=z.read(name)
            if len(content)!=meta["bytes"] or digest(content)!=meta["sha256"]:raise ValueError("ARCHIVE_HASH_MISMATCH")
        return manifest

def restore_archive(path,user):
    manifest=verify_archive(path)
    if not manifest.get("backup") or manifest.get("external"):raise ValueError("RESEARCH_BACKUP_REQUIRED")
    new_ws=uuid4()
    mapping={}
    with zipfile.ZipFile(path) as z:
        data={table:[json.loads(line) for line in z.read("data/"+table+".jsonl").splitlines()] for table in TABLES}
        asset_rows=manifest["assets"]
        for rows in [asset_rows,*data.values()]:
            for row in rows:
                if row.get("id"):mapping[row["id"]]=str(uuid4())
        mapping[manifest["workspace_id"]]=str(new_ws)
        with transaction() as conn:
            auth_ids={str(x["id"]) for x in many(conn,"select id from auth.users")}
            members=many(conn,"select * from lab.workspace_members where workspace_id=%s",(user["workspace_id"],))
            allow=many(conn,"select * from lab_private.email_allowlist where workspace_id=%s",(user["workspace_id"],))
        def remap(v,key=None):
            if key in ("created_by","reviewed_by","actor_id","reviewer_id","user_id"):
                return v if v is None or v in auth_ids else str(user["user_id"])
            if isinstance(v,str):return mapping.get(v,v)
            if isinstance(v,list):return [remap(x) for x in v]
            if isinstance(v,dict):return {mapping.get(k,k):remap(x,k) for k,x in v.items()}
            return v
        # Upload to unique new paths before linking rows; a partial failure cannot damage the active workspace.
        new_assets=[]
        store=Storage()
        for old in asset_rows:
            row=remap({k:v for k,v in old.items() if k!="archive_path"})
            row["object_path"]=str(new_ws)+"/"+row["id"]+"/"+Path(old["original_name"]).name
            if old.get("archive_path"):
                store.put(row["bucket"],row["object_path"],z.read(old["archive_path"]),row["mime_type"])
            new_assets.append(row)
        def row_types(row):
            for key,v in list(row.items()):
                if key in ("depends_on","plan_ids") and isinstance(v,list):row[key]=[UUID(x) for x in v]
                if isinstance(v,str) and key.endswith(("_at","_date")):
                    try:row[key]=datetime.fromisoformat(v)
                    except ValueError:pass
            return row
        with transaction() as conn:
            conn.execute("set constraints all deferred")
            insert(conn,"lab.workspaces",dict(id=new_ws,name="Restored TISIIF "+datetime.now().strftime("%Y-%m-%d %H:%M"),is_active=False))
            for m in members:
                insert(conn,"lab.workspace_members",{**m,"workspace_id":new_ws})
            for m in allow:
                insert(conn,"lab_private.email_allowlist",{**m,"workspace_id":new_ws})
            for row in new_assets:insert(conn,"lab.assets",row_types(row))
            new_by={r["id"]:r for r in new_assets}
            pending_accept=[]
            for table in TABLES:
                for old in data[table]:
                    row=remap(old)
                    if table=="runs":row["display_seed"]=old["display_seed"]
                    if table=="checkpoint_revisions" and row["status"]=="approved":
                        row["public_input_sha256"]=input_hash(row,new_by.get(row.get("prepared_asset_id")))
                    if table=="plans":
                        row["status"]="paused"
                        row["frozen_sha256"]=digest(row["frozen_config"])
                    if table=="jobs":
                        if row.get("accepted_attempt_id"):pending_accept.append((row["id"],row.pop("accepted_attempt_id")))
                        if row["status"] in ("dispatching","submitted","waiting_remote"):row["status"]="unknown_remote"
                        elif row["status"]=="leased":row["status"]="queued"
                        row["lease_owner"]=None;row["lease_expires_at"]=None
                        # Auxiliary export/backup jobs point to regenerable outputs not carried by the archive.
                        if row["phase"] in ("export","backup"):
                            row["status"]="cancelled";row["error_class"]="restored_regenerable_output"
                    insert(conn,"lab."+table,row_types(row))
            for job,attempt in pending_accept:
                conn.execute("update lab.jobs set accepted_attempt_id=%s where workspace_id=%s and id=%s",(attempt,new_ws,job))
            insert(conn,"lab.audit_events",dict(workspace_id=new_ws,actor_id=user["user_id"],action="workspace.restored",
                object_type="workspace",object_id=new_ws,safe_detail={"source_workspace_id":manifest["workspace_id"],
                  "uuid_mapping":mapping,"original_manifest_hash":digest(manifest),"active":False}))
    return {"workspace_id":str(new_ws),"active":False,"table_counts":manifest["table_counts"],
            "asset_count":len(new_assets),"manifest_hash":digest(manifest)}

