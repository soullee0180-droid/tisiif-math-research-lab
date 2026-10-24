import io
import json
from concurrent.futures import ThreadPoolExecutor
from uuid import UUID,uuid4
import pytest
import psycopg
from openpyxl import load_workbook
from conftest import ADMIN,OTHER,OUTSIDER,WS,token,request,seed_input,create_mock,finish_plan
from lab.db import transaction,one,many,insert
from lab.worker.engine import tick,claim,recover,process,record_failure
from lab.domain.research import digest,canonical,assemble

def test_auth_roles_and_rls(client):
    assert request(client,'/me')['role']=='admin'
    client.headers['Authorization']='Bearer '+token(OTHER,'researcher@example.test')
    assert request(client,'/me')['role']=='researcher'
    assert client.get('/api/v1/admin/health').status_code==403
    client.headers['Authorization']='Bearer '+token(OUTSIDER,'outside@example.test')
    assert client.get('/api/v1/papers').status_code==403
    client.headers.pop('Authorization')
    assert client.get('/api/v1/papers').status_code==401
    with transaction() as conn:
        conn.execute('set local role authenticated')
        conn.execute("select set_config('request.jwt.claim.sub',%s,true)",(str(OUTSIDER),))
        assert one(conn,'select count(*) n from lab.workspaces')['n']==0
        with pytest.raises(psycopg.errors.InsufficientPrivilege):conn.execute('select * from lab_private.provider_secrets')

def test_full_flow_permissions_blinding_excel_restore(client,tmp_path):
    paper,cp,rev=seed_input(client,'pdf')
    key=str(uuid4());plan=create_mock(client,rev['id'],key)
    assert create_mock(client,rev['id'],key)==plan
    assert plan['run_count']==6 and plan['logical_job_count']==36
    client.headers['Authorization']='Bearer '+token(OTHER,'researcher@example.test')
    request(client,'/me')
    assert request(client,'/plans/'+plan['plan_id'])['title']=='Fixture research flow'
    assert client.post('/api/v1/plans/'+plan['plan_id']+'/pause',json={}).status_code==403
    client.headers['Authorization']='Bearer '+token()
    finish_plan(plan['plan_id'])
    filters={'plan_id':plan['plan_id'],'transport':'mock','include_fixture':True}
    summary=request(client,'/results/summary?plan_id='+plan['plan_id']+'&transport=mock&include_fixture=true')
    assert summary['run_count']==6
    with transaction() as conn:
        assert one(conn,'select count(*) n from lab.ideas')['n']==30
        assert one(conn,"select count(*) n from lab.jobs where status='completed'")['n']==38
        assert one(conn,"select count(*) n from lab.attempts where status='accepted'")['n']==36
        assert one(conn,'select count(*) n from lab.usefulness_ratings u join lab.runs r on u.run_id=r.id where u.judge_provider=r.generator_provider')['n']==0
        gen=one(conn,"select * from lab.jobs where phase='generate' limit 1")
        attempt=one(conn,'select * from lab.attempts where id=%s',(gen['accepted_attempt_id'],))
        from lab.storage.client import Storage
        a=one(conn,'select * from lab.assets where id=%s',(attempt['request_asset_id'],))
    packet=json.loads(Storage().get(a))
    assert 'HIDDEN_CANARY' not in packet['user_text']
    batch=request(client,'/review-batches',{'plan_ids':[plan['plan_id']],'reviewer_ids':[str(ADMIN),str(OTHER)],'per_stratum':1,'seed':'synthetic'})
    assert batch['assigned_judgments']==12
    assignments=request(client,'/review-assignments')['items'];ident=assignments[0]['id']
    blinded=request(client,'/review-assignments/'+ident)
    for hidden in ['reference_idea_md','generator_provider','idea_id','stratum','selected','score']:assert hidden not in blinded
    assert client.post('/api/v1/review-assignments/'+ident+'/submit',json={'phase':'agreement','agreement_label':'core','rationale_md':'yes','expected_revision_no':0}).status_code==409
    request(client,'/review-assignments/'+ident+'/submit',{'phase':'usefulness','score':3,'validity':'promising_unverified','rationale_md':'A plausible synthetic idea','expected_revision_no':0})
    assert 'reference_idea_md' in request(client,'/review-assignments/'+ident)
    request(client,'/review-assignments/'+ident+'/submit',{'phase':'agreement','agreement_label':'partial','rationale_md':'Partial overlap','expected_revision_no':0})
    assert client.post('/api/v1/review-assignments/'+ident+'/submit',json={'phase':'agreement','agreement_label':'core','rationale_md':'Correction','expected_revision_no':1}).status_code==400
    request(client,'/review-assignments/'+ident+'/submit',{'phase':'agreement','agreement_label':'core','rationale_md':'Correction','expected_revision_no':1,'correction_reason':'Rechecked correspondence'})
    with transaction() as conn:
        assert [r['agreement_label'] for r in many(conn,"select agreement_label from lab.human_reviews where phase='agreement' order by revision_no")]==['partial','core']
    export=request(client,'/exports',{'kind':'xlsx','filters':filters,'include_source':False})
    assert tick('test-worker')
    result=request(client,'/exports/'+export['export_id']);assert result['status']=='completed'
    with transaction() as conn:a=one(conn,'select * from lab.assets where id=%s',(result['asset_id'],))
    book=load_workbook(io.BytesIO(Storage().get(a)),data_only=True)
    assert len(book.sheetnames)==10 and book['후보_아이디어'].max_row==31
    assert all(w.freeze_panes=='A2' for w in book)
    from lab.export.archive import write_archive,verify_archive,restore_archive
    actor={'workspace_id':WS,'user_id':ADMIN}
    path=tmp_path/'backup.zip';manifest=write_archive(actor,filters,path,True,backup=True)
    assert verify_archive(path)['table_counts']['ideas']==30
    bundle,parts=Storage().persist_bundle(actor,'research_backup',path,'backup.zip','application/zip',part_size=4096)
    assert parts and bundle['mime_type']=='application/json'
    joined=tmp_path/'joined.zip';Storage().download_bundle(bundle,joined)
    assert digest(joined.read_bytes())==digest(path.read_bytes())
    restored=restore_archive(joined,actor);assert not restored['active']
    with transaction() as conn:
        ws=restored['workspace_id']
        for table,count in manifest['table_counts'].items():
            actual=one(conn,'select count(*) n from lab.'+table+' where workspace_id=%s',(ws,))['n']
            assert actual==count+(1 if table=='audit_events' else 0),(table,actual,count)
        assert one(conn,'select count(*) n from lab.ideas where workspace_id=%s',(ws,))['n']==30
        assert one(conn,'select count(*) n from lab.attempts where workspace_id=%s',(ws,))['n']==36
        assert {p['status'] for p in many(conn,'select status from lab.plans where workspace_id=%s',(ws,))}=={'paused'}
        assert one(conn,'select is_active from lab.workspaces where id=%s',(WS,))['is_active']
        old=many(conn,'select display_seed from lab.runs where workspace_id=%s order by display_seed',(WS,))
        new=many(conn,'select display_seed from lab.runs where workspace_id=%s order by display_seed',(ws,))
        assert old==new
    # SQL guards remain authoritative even if an application route regresses.
    with pytest.raises(psycopg.Error):
        with transaction() as conn:conn.execute("update lab.checkpoint_revisions set goal_md='tamper' where id=%s",(rev['id'],))
    with pytest.raises(psycopg.Error):
        with transaction() as conn:conn.execute("update lab.ideas set title='tamper' where workspace_id=%s",(WS,))

def test_lease_unknown_remote_and_saved_response_no_resend(client,monkeypatch):
    _,_,rev=seed_input(client);plan=create_mock(client,rev['id'])
    with transaction() as conn:conn.execute("update lab_private.provider_runtime set max_concurrency=2 where provider='mock'")
    with ThreadPoolExecutor(max_workers=6) as executor:
        claims=list(executor.map(claim,['claim-'+str(i) for i in range(6)]))
    active=[c for c in claims if c]
    assert len(active)==2 and len({c[0]['id'] for c in active})==2
    job,attempt=active[0]
    from lab.worker.engine import reserve_response,actor_for
    with transaction() as conn:
        reserve_response(conn,actor_for(conn,job),attempt)
        conn.execute("update lab.jobs set status='dispatching',lease_expires_at=now()-interval '1 second' where id=%s",(job['id'],))
        conn.execute("update lab.attempts set status='dispatching' where id=%s",(attempt['id'],))
    recover()
    with transaction() as conn:assert one(conn,'select status from lab.jobs where id=%s',(job['id'],))['status']=='unknown_remote'
    assert client.post('/api/v1/jobs/'+str(job['id'])+'/retry',json={'reason':'Inspect'}).status_code==409
    request(client,'/jobs/'+str(job['id'])+'/retry',{'reason':'Explicit fixture retry','acknowledge_duplicate_risk':True})
    # The second claimed stage saves a response, then parsing fails. Retrying must reuse it.
    second,second_attempt=active[1]
    import lab.worker.engine as engine
    real_accept=engine.accept
    monkeypatch.setattr(engine,'accept',lambda *a,**k:(_ for _ in ()).throw(RuntimeError('synthetic crash after durable response')))
    try:process(second,second_attempt)
    except Exception as error:record_failure(second,second_attempt,error)
    with transaction() as conn:
        assert one(conn,'select status from lab.attempts where id=%s',(second_attempt['id'],))['status']=='response_saved'
        assert one(conn,'select status from lab.jobs where id=%s',(second['id'],))['status']=='needs_review'
        conn.execute("update lab.jobs set status='queued',priority=1 where id=%s",(second['id'],))
    monkeypatch.setattr(engine,'accept',real_accept)
    monkeypatch.setattr(engine,'mock_send',lambda *_:(_ for _ in ()).throw(AssertionError('MUST_NOT_SEND_TWICE')))
    picked=claim('reparse');assert picked[1]['id']==second_attempt['id']
    process(*picked)
    with transaction() as conn:assert one(conn,'select status from lab.jobs where id=%s',(second['id'],))['status']=='completed'

def test_7200_job_queue_claim_bounds(client):
    _,cp,rev=seed_input(client)
    from lab.domain.plans import create_plan
    from lab.api.types import PlanBody
    with transaction() as conn:
        source=one(conn,'select * from lab.checkpoint_revisions where id=%s',(rev['id'],))
        revisions=[source['id']]
        for i in range(2,201):
            row={**source,'id':uuid4(),'version_no':i};insert(conn,'lab.checkpoint_revisions',row);revisions.append(row['id'])
        protocol=one(conn,'select id from lab.protocol_versions where workspace_id=%s',(WS,))
        body=PlanBody(title='7200 logical-job synthetic queue',cohort='fixture',transport='mock',checkpoint_revision_ids=revisions,protocol_id=protocol['id'])
        result=create_plan(conn,{'workspace_id':WS,'user_id':ADMIN},body,uuid4())
        assert result['logical_job_count']==7200
        conn.execute("update lab_private.provider_runtime set max_concurrency=4 where provider='mock'")
    with ThreadPoolExecutor(max_workers=12) as executor:
        results=list(executor.map(claim,['stress-'+str(i) for i in range(24)]))
    active=[r for r in results if r]
    assert len(active)<=4
    while len(active)<4:
        extra=claim('capacity-check');assert extra is not None;active.append(extra)
    assert claim('over-capacity') is None
    assert len(active)==4 and len({r[0]['id'] for r in active})==4
    with transaction() as conn:
        assert one(conn,"select count(*) n from lab.jobs where plan_id=%s",(result['plan_id'],))['n']==7200
        assert one(conn,"select count(*) n from lab.jobs where plan_id=%s and status='leased'",(result['plan_id'],))['n']==4

def test_manual_transport_never_dispatches_and_keeps_original(client):
    _,_,rev=seed_input(client)
    profiles={}
    for provider in ('google','openai','anthropic'):
        p=request(client,'/admin/model-profiles',{'provider':provider,'label':'Synthetic manual profile','model_id':'fixture-ui-'+provider,
            'config':{},'capabilities':{'pdf_supported':True,'structured_output_supported':True,'official_source_url':'https://example.test/synthetic-fixture'},'manual_only':True})
        profiles[provider]=p['id']
    protocol=request(client,'/protocols')['items'][0]
    body={'title':'Manual synthetic flow','transport':'manual','cohort':'manual','checkpoint_revision_ids':[rev['id']],'protocol_id':protocol['id'],'generator_profile_ids':profiles}
    plan=request(client,'/plans',body)
    assert claim('must-not-dispatch-manual') is None
    with transaction() as conn:job=one(conn,"select * from lab.jobs where plan_id=%s and phase='generate' order by id limit 1",(plan['plan_id'],))
    ident=str(job['id']);prompt=request(client,'/jobs/'+ident+'/manual-prompt')
    assert 'HIDDEN_CANARY' not in prompt['prompt']
    from lab.providers.adapters import mock_send
    raw=mock_send(prompt['packet']).text
    response={'raw_text':raw,'screen_model':'Synthetic manual UI','mode':'fresh conversation','occurred_at':'2026-10-24T00:00:00+00:00','tools':'off','web_search':'off','memory':'off'}
    key=str(uuid4());accepted=request(client,'/jobs/'+ident+'/manual-response',response,key=key)
    assert accepted['accepted']
    assert request(client,'/jobs/'+ident+'/manual-response',response,key=key)==accepted
    recover()
    with transaction() as conn:selection=one(conn,"select * from lab.jobs where run_id=%s and phase='select'",(job['run_id'],))
    target=str(selection['id']);packet=request(client,'/jobs/'+target+'/manual-prompt')['packet']
    unstructured={**response,'raw_text':'Synthetic free-form response that is deliberately not JSON.'}
    assert not request(client,'/jobs/'+target+'/manual-response',unstructured)['accepted']
    mapped={**unstructured,'mapped_response':json.loads(mock_send(packet).text),'correction_reason':'Manual schema mapping for synthetic test'}
    result=request(client,'/jobs/'+target+'/manual-response',mapped);assert result['accepted']
    with transaction() as conn:
        attempts=many(conn,'select * from lab.attempts where job_id=%s order by attempt_no',(selection['id'],))
        assert len(attempts)==2 and attempts[0]['status']=='invalid' and attempts[1]['status']=='accepted'
        assert attempts[1]['normalized_settings']['human_corrected']
        asset=one(conn,'select * from lab.assets where id=%s',(attempts[1]['raw_text_asset_id'],))
    from lab.storage.client import Storage
    assert Storage().get(asset).decode()==unstructured['raw_text']
